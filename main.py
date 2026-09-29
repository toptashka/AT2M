from typing import Optional
from pydantic import BaseModel
from fastapi import FastAPI, Depends, HTTPException, Request, UploadFile, File, Form
from sqlalchemy.orm import Session
from sqlalchemy import func
import boto3
from botocore.exceptions import ClientError
import uuid
import logging
import openpyxl
from io import BytesIO
import models
import workflow_service as workflow
from schema_upgrade import upgrade_schema
from security import authenticate, require_role
from pdf_reports import partnership_report
from catalog_import import inspect_workbook, import_rows
from database import engine, get_db
from fastapi.responses import StreamingResponse
from datetime import datetime, date, timedelta
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os
import json
import urllib.request
import urllib.error
import urllib.parse
from fastapi_cache import FastAPICache
from fastapi_cache.backends.inmemory import InMemoryBackend
from fastapi_cache.decorator import cache
from urllib.parse import quote

models.Base.metadata.create_all(bind=engine)
upgrade_schema(engine)

app = FastAPI(
    title="RTK CRM API",
    description="API для системы управления партнерствами вузов",
    version="1.0.0"
)

DEFAULT_STAGES = [
    (1, "Поиск контактов"),
    (2, "Коммуникация с вузом"),
    (3, "Встреча с представителями"),
    (4, "Обмен документами"),
    (5, "Корректировка документов"),
    (6, "Подписание документов"),
    (7, "Материалы, лицензия и документы"),
    (8, "Сопровождение внедрения"),
    (9, "Обучение преподавателей"),
    (10, "Актуализация учебной программы"),
    (11, "Ведение занятий"),
    (12, "Актуализация документации"),
    (13, "Повышение квалификации"),
    (14, "Контроль исполнения этапов"),
]

def ensure_workflow_stages():
    db = next(get_db())
    try:
        count = db.query(models.WorkflowStage).count()
        if count == 0:
            for step_num, title in DEFAULT_STAGES:
                stage = models.WorkflowStage(step_number=step_num, title=title)
                db.add(stage)
            db.commit()
    except Exception as e:
        db.rollback()
        print(f"Инициализация этапов воронки: {e}")
    finally:
        db.close()

S3_BUCKET_NAME = "rtk-crm-documents"
s3_client = boto3.client(
    's3',
    endpoint_url=os.getenv("MINIO_ENDPOINT", "http://minio:9000"),
    aws_access_key_id=os.getenv("MINIO_ROOT_USER", "admin"),
    aws_secret_access_key=os.getenv("MINIO_ROOT_PASSWORD", "admin_password")
)

def init_s3():
    try:
        s3_client.head_bucket(Bucket=S3_BUCKET_NAME)
    except Exception:
        try:
            s3_client.create_bucket(Bucket=S3_BUCKET_NAME)
        except Exception as e:
            print(f"MinIO инициализация: {e}")

@app.on_event("startup")
def startup():
    FastAPICache.init(InMemoryBackend())
    ensure_workflow_stages()
    db = next(get_db())
    try:
        workflow.lock_workflow(db)
        db.commit()
    finally:
        db.close()
    init_s3()

class PartnershipCreate(BaseModel):
    university_name: str
    program_name: str
    manager_name: Optional[str] = None
    stage_number: Optional[int] = 1


def log_audit(db: Session, request: Request, user_id: str, action: str, entity_name: str, entity_id: int):
    ip_address = request.client.host if request.client else "unknown"
    audit_entry = models.AuditLog(
        user_id=user_id,
        action=action,
        entity_name=entity_name,
        entity_id=entity_id,
        ip_address=ip_address
    )
    db.add(audit_entry)
    db.commit()

def fetch_keycloak_directory():
    kc_host = os.getenv("KEYCLOAK_URL", "http://keycloak:8080")
    admin_user = os.getenv("KEYCLOAK_ADMIN", os.getenv("KC_ADMIN", "admin"))
    realm = os.getenv("KEYCLOAK_REALM", "rtk_crm")

    passwords_to_try = [
        os.getenv("KEYCLOAK_ADMIN_PASSWORD"),
        os.getenv("KC_ADMIN_PASSWORD"),
    ]
    passwords_to_try = list(dict.fromkeys([p for p in passwords_to_try if p]))

    token = None
    prefix = ""

    for pwd in passwords_to_try:
        data = urllib.parse.urlencode({
            "client_id": "admin-cli",
            "username": admin_user,
            "password": pwd,
            "grant_type": "password"
        }).encode("utf-8")

        for path in ["/realms/master/protocol/openid-connect/token", "/auth/realms/master/protocol/openid-connect/token"]:
            try:
                req = urllib.request.Request(f"{kc_host}{path}", data=data, method="POST")
                req.add_header("Content-Type", "application/x-www-form-urlencoded")
                with urllib.request.urlopen(req, timeout=3) as resp:
                    token = json.loads(resp.read().decode("utf-8")).get("access_token")
                    prefix = path.replace("/realms/master/protocol/openid-connect/token", "")
                    if token:
                        break
            except Exception:
                continue
        if token:
            break

    if not token:
        raise RuntimeError("Keycloak Admin API unavailable")

    def read_json(url):
        req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req, timeout=3) as response:
            return json.loads(response.read().decode("utf-8"))
    base = f"{kc_host}{prefix}/admin/realms/{realm}"
    clients = read_json(base + "/clients?clientId=" + urllib.parse.quote(os.getenv("KC_CLIENT_ID", "frontend-app")))
    client_ids = [client["id"] for client in clients if client.get("clientId") == os.getenv("KC_CLIENT_ID", "frontend-app")]
    def role_members(role):
        result = {}
        paths = [base + "/roles/"] + [base + "/clients/" + id + "/roles/" for id in client_ids]
        for path in paths:
            offset = 0
            while True:
                try:
                    batch = read_json(path + urllib.parse.quote(role) + f"/users?first={offset}&max=100")
                except urllib.error.HTTPError as error:
                    if error.code == 404:
                        break
                    raise
                result.update({user["id"]: user for user in batch})
                if len(batch) < 100:
                    break
                offset += len(batch)
        return list(result.values())
    users = role_members("Пользователь")
    excluded = {u["id"] for role in ("Руководитель", "Администратор") for u in role_members(role)}
    return [{"username": u["username"], "name": " ".join(filter(None, [u.get("firstName", "").strip(), u.get("lastName", "").strip()])) or u["username"]}
            for u in users if u.get("enabled", True) and u.get("username") and u["id"] not in excluded]


_directory = {"at": 0, "rows": []}


def directory():
    import time
    if time.monotonic() - _directory["at"] > 60:
        rows = fetch_keycloak_directory()
        _directory.update(at=time.monotonic(), rows=rows)
    return _directory["rows"]


def get_keycloak_users():
    return [u["username"] for u in directory()]


def display_name(username, user=None):
    if user is not None and username == str(user) and user.display_name != str(user):
        return user.display_name
    try:
        return next((u["name"] for u in directory() if u["username"] == username), username or "")
    except Exception:
        return username or ""


def card_detail(db, row, user):
    result = workflow.detail(db, row)
    result["manager_display_name"] = display_name(row.manager_name, user)
    for stage in result["stages"].values():
        for comment in stage["comments"]:
            comment["authorId"] = comment["author"]
            comment["author"] = display_name(comment["author"], user)
    return result


@app.get("/api/v1/staff")
def staff(current_user=Depends(require_role("Пользователь"))):
    try:
        rows = directory()
    except Exception:
        raise HTTPException(503, "Не удалось загрузить имена сотрудников из Keycloak")
    return rows


@app.get("/api/v1/me")
def me(current_user=Depends(authenticate)):
    return {"username": str(current_user), "name": display_name(str(current_user), current_user), "roles": sorted(current_user.roles),
            "canManage": current_user.manager, "canAdmin": current_user.admin}


@app.get("/api/v1/managers", response_model=list[str])
def get_managers(db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    if not current_user.manager:
        return [str(current_user)]
    try:
        users = get_keycloak_users()
    except Exception:
        users = []
    return sorted(set(users))


def validate_owner(name, current_user):
    if not current_user.manager:
        if name not in (None, "", str(current_user)):
            raise HTTPException(403, "Назначать ответственных может только руководитель или администратор")
        return str(current_user)
    if not name:
        return ""
    try:
        users = get_keycloak_users()
    except Exception:
        raise HTTPException(503, "Не удалось проверить пользователя в Keycloak")
    if name not in users:
        raise HTTPException(422, "Выберите логин пользователя Keycloak, а не произвольное ФИО")
    return name


@app.get("/api/v1/partnerships")
def get_partnerships_grid(stage_id: Optional[int] = None, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    query = workflow.scope(db, current_user)
    if stage_id:
        query = query.join(models.WorkflowStage).filter(models.WorkflowStage.step_number == stage_id)
    return [card_detail(db, partnership, current_user) for partnership in query.order_by(models.Partnership.id).all()]


@app.post("/api/v1/partnerships")
def create_partnership(data: PartnershipCreate, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    workflow.lock_workflow(db)
    uni = db.query(models.University).filter_by(name=data.university_name.strip()).first()
    prog = db.query(models.Program).filter(
        (models.Program.direction == data.program_name.strip()) | (models.Program.name == data.program_name.strip())
    ).order_by(models.Program.id).first()
    if not uni or not prog:
        raise HTTPException(422, "Выберите учреждение и направление из загруженных справочников")
    stage = db.query(models.WorkflowStage).order_by(models.WorkflowStage.step_number).first()
    if not stage:
        raise HTTPException(422, "Не настроен Workflow")
    if not current_user.manager:
        if data.manager_name not in (None, "", str(current_user)):
            raise HTTPException(403, "Ответственного назначает руководитель при одобрении заявки")
        pending = models.PartnershipRequest(requester=str(current_user), requester_name=display_name(str(current_user), current_user),
                                            university_id=uni.id, program_id=prog.id)
        db.add(pending)
        db.flush()
        workflow.audit(db, current_user, request, "CREATE_REQUEST", "partnership_requests", pending.id)
        db.commit()
        return {"kind": "request", "request": request_detail(db, pending)}
    owner = validate_owner(data.manager_name, current_user)
    row = models.Partnership(university_id=uni.id, program_id=prog.id, stage_id=stage.id, manager_name=owner)
    db.add(row)
    db.flush()
    workflow.initialize_deadline(db, row)
    workflow.audit(db, current_user, request, "CREATE_PARTNERSHIP", "partnerships", row.id)
    db.commit()
    return card_detail(db, row, current_user)


def request_detail(db, row):
    uni = db.get(models.University, row.university_id)
    program = db.get(models.Program, row.program_id)
    return {"id": row.id, "name": uni.name, "program": program.direction or program.name,
            "initiator": row.requester_name, "requester": row.requester, "source": "Пользователь",
            "createdAt": row.created_at.isoformat(), "time": row.created_at.strftime("%d.%m.%Y"),
            "status": row.status, "reason": row.reason or "", "partnershipId": row.partnership_id, "details": []}


@app.get("/api/v1/requests")
def list_requests(db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    query = db.query(models.PartnershipRequest)
    if not current_user.manager:
        query = query.filter_by(requester=str(current_user))
    return [request_detail(db, row) for row in query.order_by(models.PartnershipRequest.created_at.desc()).all()]


@app.post("/api/v1/requests/{request_id}/decision")
def decide_request(request_id: int, payload: dict, request: Request, db: Session = Depends(get_db),
                   current_user=Depends(require_role("Руководитель"))):
    workflow.lock_workflow(db)
    pending = db.query(models.PartnershipRequest).filter_by(id=request_id).with_for_update().first()
    if pending is None:
        raise HTTPException(404, "Заявка не найдена")
    if pending.status != "pending":
        raise HTTPException(409, "Заявка уже рассмотрена. Обновите страницу.")
    if payload.get("decision") not in {"approve", "reject"}:
        raise HTTPException(422, "Выберите одобрение или отклонение")
    reason = payload.get("comment", "")
    if not isinstance(reason, str) or len(reason) > 10000:
        raise HTTPException(422, "Некорректный комментарий")
    if payload["decision"] == "approve":
        owner = validate_owner(payload.get("manager_name"), current_user)
        if not owner:
            raise HTTPException(422, "Выберите ответственного")
        first = db.query(models.WorkflowStage).order_by(models.WorkflowStage.step_number).first()
        if not first:
            raise HTTPException(422, "Не настроен Workflow")
        row = models.Partnership(university_id=pending.university_id, program_id=pending.program_id,
                                 stage_id=first.id, manager_name=owner)
        db.add(row)
        db.flush()
        workflow.initialize_deadline(db, row)
        pending.partnership_id = row.id
        pending.status = "approved"
    else:
        pending.status = "rejected"
    pending.reason = reason.strip()
    pending.decided_by = str(current_user)
    pending.decided_at = datetime.utcnow()
    workflow.audit(db, current_user, request, "REQUEST_" + pending.status.upper(), "partnership_requests", pending.id)
    db.commit()
    return request_detail(db, pending)


@app.get("/api/v1/partnerships/{partnership_id}")
def get_partnership_detail(partnership_id: int, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    row = workflow.accessible(db, current_user, partnership_id)
    result = card_detail(db, row, current_user)
    workflow.audit(db, current_user, request, "VIEW_PDN", "partnerships", row.id)
    db.commit()
    return result


def apply_card_fields(row, state, payload):
    if "contract" in payload:
        contract = payload["contract"]
        if not isinstance(contract, dict):
            raise HTTPException(422, "Некорректные реквизиты")
        allowed = {"vendor", "software", "number", "licenseEnd", "signed", "transfer"}
        if set(contract) - allowed or any(not isinstance(v, str) or len(v) > 2000 for v in contract.values()):
            raise HTTPException(422, "Некорректные реквизиты")
        if contract.get("signed", "Нет") not in {"Нет", "В процессе", "Да", "Да / Подписан"}:
            raise HTTPException(422, "Некорректный статус подписания")
        license_end = contract.get("licenseEnd", "")
        if license_end:
            try:
                date.fromisoformat(license_end)
            except ValueError:
                raise HTTPException(422, "Срок лицензии должен быть датой ГГГГ-ММ-ДД")
        state.setdefault("contract", {}).update(contract)
        if "number" in contract:
            if len(contract["number"]) > 100:
                raise HTTPException(422, "Номер договора не должен превышать 100 символов")
            row.contract_number = contract["number"]
        if "signed" in contract:
            row.is_license_signed = contract["signed"] in {"Да", "Да / Подписан"}
        if "transfer" in contract:
            if len(contract["transfer"]) > 100:
                raise HTTPException(422, "Слишком длинный статус передачи")
            row.transfer_status = contract["transfer"]
    if "contact" in payload:
        contact = payload["contact"]
        if not isinstance(contact, dict) or set(contact) - {"name", "position", "phone", "email"} or any(not isinstance(v, str) or len(v) > 500 for v in contact.values()):
            raise HTTPException(422, "Некорректные контактные данные")
        state.setdefault("contact", {}).update(contact)


@app.patch("/api/v1/partnerships/{partnership_id}")
def update_partnership(partnership_id: int, payload: dict, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    workflow.lock_workflow(db)
    row = workflow.accessible(db, current_user, partnership_id, lock=True)
    workflow.check_current(db, row, payload.get("expected_stage_id"), payload.get("workflow_version"))
    if "manager_name" in payload:
        if not current_user.manager:
            raise HTTPException(403, "Назначать ответственного может только руководитель или администратор")
        row.manager_name = validate_owner(payload["manager_name"], current_user)
    state = workflow.load_state(db, row.id)
    apply_card_fields(row, state, payload)
    row.updated_at = datetime.utcnow()
    workflow.save_state(db, row.id, state)
    workflow.audit(db, current_user, request, "UPDATE_PARTNERSHIP", "partnerships", row.id)
    db.commit()
    return card_detail(db, row, current_user)


@app.patch("/api/v1/partnerships/{partnership_id}/conditions")
def update_conditions(partnership_id: int, payload: dict, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    workflow.lock_workflow(db)
    row = workflow.accessible(db, current_user, partnership_id, lock=True)
    workflow.check_current(db, row, payload.get("expected_stage_id"), payload.get("workflow_version"))
    index = payload.get("index")
    value = payload.get("value")
    if type(index) is not int or index not in {0, 2} or type(value) is not bool:
        raise HTTPException(422, "Условие прикрепления файла определяется сервером")
    state = workflow.load_state(db, row.id)
    if state.get("completed"):
        raise HTTPException(409, "Взаимодействие завершено")
    stage = state.setdefault("stages", {}).setdefault(str(row.stage_id), {})
    values = stage.setdefault("conditions", [False, False, False])
    values[index] = value
    workflow.save_state(db, row.id, state)
    row.updated_at = datetime.utcnow()
    workflow.audit(db, current_user, request, "UPDATE_CONDITION", "partnerships", row.id)
    db.commit()
    return card_detail(db, row, current_user)


@app.patch("/api/v1/partnerships/{partnership_id}/stage")
def update_partnership_stage(partnership_id: int, payload: dict, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    workflow.lock_workflow(db)
    row = workflow.accessible(db, current_user, partnership_id, lock=True)
    workflow.advance(db, current_user, request, row, payload)
    return card_detail(db, row, current_user)


@app.get("/api/v1/partnerships/{partnership_id}/comments")
def get_partnership_comments(partnership_id: int, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    row = workflow.accessible(db, current_user, partnership_id)
    return [c for stage in card_detail(db, row, current_user)["stages"].values() for c in stage["comments"]]


@app.post("/api/v1/partnerships/{partnership_id}/comments")
def add_partnership_comment(partnership_id: int, payload: dict, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    workflow.lock_workflow(db)
    row = workflow.accessible(db, current_user, partnership_id, lock=True)
    workflow.check_current(db, row, payload.get("expected_stage_id"), payload.get("workflow_version"))
    value = payload.get("text")
    if not isinstance(value, str) or not value.strip() or len(value) > 10000:
        raise HTTPException(422, "Введите комментарий до 10000 символов")
    if workflow.load_state(db, row.id).get("completed"):
        raise HTTPException(409, "Взаимодействие завершено")
    db.add(models.PartnershipComment(partnership_id=row.id, stage_id=row.stage_id, author_id=str(current_user), text=value.strip()))
    row.updated_at = datetime.utcnow()
    workflow.audit(db, current_user, request, "ADD_COMMENT", "partnerships", row.id)
    db.commit()
    return card_detail(db, row, current_user)


@app.post("/api/v1/partnerships/{partnership_id}/files")
def upload_partnership_file(partnership_id: int, request: Request, file: UploadFile = File(...),
                            document_type: str = Form("other"), expected_stage_id: int = Form(...), workflow_version: int = Form(...),
                            metadata: str = Form("{}"), comment: str = Form(""),
                            db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    workflow.lock_workflow(db)
    row = workflow.accessible(db, current_user, partnership_id, lock=True)
    workflow.check_current(db, row, expected_stage_id, workflow_version)
    if workflow.load_state(db, row.id).get("completed"):
        raise HTTPException(409, "Взаимодействие завершено")
    name = (file.filename or "").replace("\\", "/").split("/")[-1]
    extension = name.rsplit(".", 1)[-1].lower()
    if not name or len(name) > 255 or extension not in {"png", "jpg", "jpeg", "pdf", "zip", "gzip", "gz", "rar", "doc", "docx", "xls", "xlsx"}:
        raise HTTPException(422, "Неподдерживаемое имя или формат файла")
    file.file.seek(0, 2)
    size = file.file.tell()
    file.file.seek(0)
    if not size or size > 50 * 1024 * 1024:
        raise HTTPException(422, "Размер файла должен быть от 1 байта до 50 МБ")
    if document_type not in {"license", "contract", "materials", "other"}:
        raise HTTPException(422, "Некорректный тип документа")
    state = workflow.load_state(db, row.id)
    try:
        contract = json.loads(metadata)
    except ValueError:
        raise HTTPException(422, "Некорректные метаданные файла")
    apply_card_fields(row, state, {"contract": contract})
    if len(comment) > 10000:
        raise HTTPException(422, "Комментарий слишком длинный")
    key = f"partnership_{row.id}/{uuid.uuid4()}.{extension}"
    try:
        s3_client.upload_fileobj(file.file, S3_BUCKET_NAME, key, ExtraArgs={"ContentType": file.content_type or "application/octet-stream"})
    except Exception:
        raise HTTPException(502, "MinIO недоступен. Файл не прикреплён.")
    try:
        attachment = models.Attachment(partnership_id=row.id, stage_id=row.stage_id, file_name=name, file_url=key,
                                       file_size=size, document_type=document_type, file_type=file.content_type)
        db.add(attachment)
        workflow.save_state(db, row.id, state)
        if comment.strip():
            db.add(models.PartnershipComment(partnership_id=row.id, stage_id=row.stage_id, author_id=str(current_user), text=comment.strip()))
        row.updated_at = datetime.utcnow()
        workflow.audit(db, current_user, request, "UPLOAD_FILE", "partnerships", row.id)
        db.commit()
    except Exception:
        db.rollback()
        try:
            s3_client.delete_object(Bucket=S3_BUCKET_NAME, Key=key)
        except Exception:
            logging.getLogger(__name__).exception("Failed to remove uncommitted upload")
        raise HTTPException(500, "Не удалось сохранить сведения о файле")
    return card_detail(db, row, current_user)


@app.get("/api/v1/files/{file_id}")
def download_partnership_file(file_id: int, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    attachment = db.get(models.Attachment, file_id)
    if attachment is None:
        raise HTTPException(404, "Файл не найден")
    workflow.accessible(db, current_user, attachment.partnership_id)
    try:
        response = s3_client.get_object(Bucket=S3_BUCKET_NAME, Key=attachment.file_url)
    except Exception:
        raise HTTPException(502, "Файл недоступен в MinIO")
    workflow.audit(db, current_user, request, "DOWNLOAD_FILE", "attachments", attachment.id)
    db.commit()
    def stream():
        try:
            while True:
                chunk = response["Body"].read(65536)
                if not chunk:
                    break
                yield chunk
        finally:
            response["Body"].close()
    return StreamingResponse(stream(), media_type=response.get("ContentType", "application/octet-stream"),
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(attachment.file_name)}"})


@app.delete("/api/v1/files/{file_id}")
def delete_partnership_file(file_id: int, request: Request, expected_stage_id: int, workflow_version: int,
                            db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    workflow.lock_workflow(db)
    attachment = db.get(models.Attachment, file_id)
    if attachment is None:
        raise HTTPException(404, "Файл не найден")
    row = workflow.accessible(db, current_user, attachment.partnership_id, lock=True)
    workflow.check_current(db, row, expected_stage_id, workflow_version)
    if (attachment.stage_id or row.stage_id) != row.stage_id or workflow.load_state(db, row.id).get("completed"):
        raise HTTPException(409, "Файлы завершённых этапов доступны только для чтения")
    try:
        s3_client.delete_object(Bucket=S3_BUCKET_NAME, Key=attachment.file_url)
    except Exception:
        raise HTTPException(502, "MinIO недоступен. Удаление не подтверждено.")
    db.delete(attachment)
    row.updated_at = datetime.utcnow()
    workflow.audit(db, current_user, request, "DELETE_FILE", "attachments", file_id)
    db.commit()
    return card_detail(db, row, current_user)


@app.post("/api/v1/partnerships/{partnership_id}/students/import")
def import_students_from_excel(
    partnership_id: int,
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Пользователь"))
):
    partnership = workflow.accessible(db, current_user, partnership_id)

    if not file.filename.endswith(('.xls', '.xlsx')):
        raise HTTPException(status_code=400, detail="Поддерживаются только форматы xls и xlsx")

    try:
        contents = file.file.read()
        workbook = openpyxl.load_workbook(filename=BytesIO(contents), data_only=True)
        sheet = workbook.active
        students_added = 0
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                new_student = models.Student(
                    partnership_id=partnership.id,
                    full_name=str(row[0]).strip(),
                    email=str(row[1]).strip()
                )
                db.add(new_student)
                students_added += 1
                        
        db.commit()
        log_audit(db, request, user_id=current_user, action=f"IMPORT_STUDENTS_COUNT_{students_added}", entity_name="partnerships", entity_id=partnership.id)
        return {"status": "ok", "message": f"Успешно загружено студентов: {students_added}"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка обработки файла: {str(e)}")

@app.post("/api/v1/catalogs/inspect")
async def inspect_catalogs(
    request: Request,
    file: UploadFile = File(...),
    catalog: str = Form("institutions"),
    partnershipId: Optional[int] = Form(None, alias="partnership_id"),
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Администратор"))
):
    inspected = inspect_workbook(await file.read(), file.filename or "", catalog)
    if catalog == "students" and not db.get(models.Partnership, partnershipId):
        raise HTTPException(422, "Выберите существующее партнёрство")
    inspected["partnershipId"] = partnershipId
    inspected["filename"] = file.filename
    job_id = str(uuid.uuid4())
    try:
        db.query(models.CatalogImportJob).filter(
            models.CatalogImportJob.created_at < datetime.utcnow() - timedelta(days=2)
        ).delete(synchronize_session=False)
        db.add(models.CatalogImportJob(
            id=job_id, owner_id=current_user, catalog=catalog,
            payload=json.dumps(inspected, ensure_ascii=False)
        ))
        db.commit()
    except Exception:
        db.rollback()
        raise
    columns = [{"id": h, "label": h} for h in inspected["headers"]]
    rows = [{"id": f"row-{i + 1}", "values": row} for i, row in enumerate(inspected["rows"])]
    return {
        "jobId": job_id, "columns": columns, "rows": rows,
        "preview": rows[:20], "total": len(rows), "mapping": inspected["mapping"],
        "filename": file.filename
    }


@app.post("/api/v1/catalogs/import")
async def import_catalogs(
    request: Request,
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Администратор"))
):
    try:
        body = await request.json()
        if not isinstance(body, dict) or not isinstance(body.get("jobId"), str):
            raise HTTPException(422, "Сначала загрузите XLSX-файл для проверки")
        job = db.query(models.CatalogImportJob).filter_by(
            id=body["jobId"], owner_id=current_user
        ).with_for_update().first()
        if job is None or job.created_at < datetime.utcnow() - timedelta(days=2):
            raise HTTPException(410, "Подготовленная загрузка не найдена или устарела. Загрузите файл повторно.")
        if body.get("catalog", job.catalog) != job.catalog:
            raise HTTPException(422, "Тип справочника отличается от проверенного файла")
        if job.result:
            return json.loads(job.result)
        try:
            data = json.loads(job.payload)
        except (ValueError, TypeError):
            raise HTTPException(409, "Не удалось прочитать загрузку. Проверьте постоянный ENCRYPTION_KEY и загрузите файл повторно.")
        partnership_id = data.get("partnershipId")
        if job.catalog == "students" and str(body.get("partnershipId")) != str(partnership_id):
            raise HTTPException(422, "Партнёрство отличается от выбранного при загрузке")
        result = import_rows(
            db, job.catalog, data["rows"], body.get("mapping", data["mapping"]),
            body.get("excludedRowIds", []), partnership_id
        )
        job.result = json.dumps(result, ensure_ascii=False)
        db.add(models.AuditLog(
            user_id=current_user, action=f"IMPORT_CATALOGS_PROCESSED_{result['processed']}_ADDED_{result['added']}",
            entity_name="catalogs", entity_id=0,
            ip_address=request.client.host if request.client else "unknown"
        ))
        db.commit()
        return result
    except HTTPException:
        db.rollback()
        raise
    except Exception as exc:
        db.rollback()
        logging.getLogger(__name__).exception("Catalog import failed")
        raise HTTPException(500, "Импорт не выполнен, изменения отменены. Проверьте журнал API.") from exc


class CMSLead(BaseModel):
    university_name: str
    contact_name: str
    contact_email: str
    program_name: str

class LMSSyncData(BaseModel):
    partnership_id: int
    active_students: int
    average_score: float

class ProgramResponse(BaseModel):
    id: int
    name: str
    direction: Optional[str] = None
    vendor: Optional[str] = None
    software: Optional[str] = None
    priority: int

class ProgramPriorityUpdate(BaseModel):
    priority: int

@app.post("/api/v1/integrations/cms/leads")
def receive_lead_from_cms(lead: CMSLead, current_user=Depends(require_role("Администратор"))):
    raise HTTPException(501, "Интеграция CMS не настроена")


@app.post("/api/v1/integrations/lms/sync")
def sync_with_lms(data: LMSSyncData, current_user=Depends(require_role("Администратор"))):
    raise HTTPException(501, "Интеграция LMS не настроена")


@app.get("/api/v1/reports/partnerships/excel")
def export_partnerships_excel(
    request: Request,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Пользователь"))
):
    query = workflow.scope(db, current_user)
    if start_date:
        query = query.filter(models.Partnership.created_at >= start_date)
    if end_date:
        query = query.filter(models.Partnership.created_at <= end_date)
    partnerships = query.all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчет по взаимодействиям"
    
    headers = ["ID", "ВУЗ", "ИТ-Программа", "Текущий статус", "Ответственный РТК", "Номер договора", "Лицензия подписана"]
    ws.append(headers)
    
    for p in partnerships:
        direction_name = p.program.direction if (p.program and p.program.direction) else (p.program.name if p.program else "—")
        ws.append([
            p.id,
            p.university.name if p.university else "—",
            direction_name,
            p.stage.title if p.stage else "—",
            display_name(p.manager_name, current_user) or "Не назначен",
            p.contract_number or "Нет данных",
            "Да" if p.is_license_signed else "Нет"
        ])
            
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    filename = f"rtk_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.get("/api/v1/reports/partnerships/pdf")
def export_partnerships_pdf(
    request: Request,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Пользователь"))
):
    query = workflow.scope(db, current_user)
    if start_date:
        query = query.filter(models.Partnership.created_at >= start_date)
    if end_date:
        query = query.filter(models.Partnership.created_at <= end_date)
    partnerships = query.all()
    
    rows = []
    for p in partnerships:
        rows.append([p.id, p.university.name, p.program.direction or p.program.name, p.stage.title,
                     display_name(p.manager_name, current_user), p.contract_number, "Да" if p.is_license_signed else "Нет"])
    stream = partnership_report(rows)

    filename = f"rtk_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return StreamingResponse(
        stream,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.get("/api/v1/programs", response_model=list[ProgramResponse])
def get_programs_catalog(
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Пользователь"))
):
    programs = db.query(models.Program).order_by(models.Program.priority.desc()).all()
    return [
        ProgramResponse(
            id=p.id,
            name=p.name,
            direction=p.direction,
            vendor=p.vendor,
            software=p.software,
            priority=p.priority
        ) for p in programs
    ]

@app.patch("/api/v1/programs/{program_id}/priority", response_model=ProgramResponse)
def update_program_priority(
    program_id: int,
    priority_data: ProgramPriorityUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Руководитель"))
):
    program = db.query(models.Program).filter(models.Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="ИТ-программа не найдена")
            
    program.priority = priority_data.priority
    db.commit()
    db.refresh(program)
    return ProgramResponse(id=program.id, name=program.name, direction=program.direction, vendor=program.vendor, software=program.software, priority=program.priority)

class StageResponse(BaseModel):
    id: int
    step_number: int
    title: str
    description: Optional[str] = None

@app.get("/api/v1/workflow/stages")
def get_workflow_stages(db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    state = workflow.workflow_payload(db)
    return [{**stage, "version": state["version"]} for stage in state["stages"]]


@app.patch("/api/v1/workflow/deadlines")
def update_stage_deadlines(payload: dict, request: Request, db: Session = Depends(get_db),
                           current_user=Depends(require_role("Руководитель"))):
    try:
        return workflow.update_deadlines(db, current_user, request, payload)
    except Exception:
        db.rollback()
        raise


@app.get("/api/v1/workflow")
def get_workflow(db: Session = Depends(get_db), current_user=Depends(require_role("Пользователь"))):
    return workflow.workflow_payload(db)


@app.post("/api/v1/workflow/stages")
def update_workflow(payload: dict, request: Request, db: Session = Depends(get_db), current_user=Depends(require_role("Администратор"))):
    try:
        return workflow.modify_workflow(db, current_user, request, payload)
    except Exception:
        db.rollback()
        raise


class AuditLogResponse(BaseModel):
    id: str
    date: str
    user: str
    action: str
    object: str
    ip: str

@app.get("/api/v1/audit", response_model=list[AuditLogResponse])
def get_audit_logs(
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Администратор"))
):
    logs = db.query(models.AuditLog).order_by(models.AuditLog.timestamp.desc()).all()
    return [
        AuditLogResponse(
            id=str(log.id),
            date=log.timestamp.isoformat() if log.timestamp else "",
            user=log.user_id or "Система",
            action=log.action,
            object=f"{log.entity_name} #{log.entity_id}" if log.entity_id else (log.entity_name or "—"),
            ip=log.ip_address or "—",
        ) for log in logs
    ]

@app.get("/api/v1/universities", response_model=list[str])
def get_universities_catalog(
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Пользователь"))
):
    return [u.name for u in db.query(models.University).order_by(models.University.name).all() if u.name]

@app.get("/")
def read_root():
    return {"status": "ok", "message": "Бэкенд успешно запущен!"}

@app.get("/api/v1/catalogs/options")
def get_catalog_options(
    db: Session = Depends(get_db),
    current_user: str = Depends(require_role("Пользователь"))
):
    universities = db.query(models.University).order_by(models.University.name).all()
    return {
        "programs": get_programs_catalog(db, current_user),
        "universities": [u.name for u in universities],
        "regions": sorted({u.region for u in universities if u.region}),
    }
