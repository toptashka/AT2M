import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

os.makedirs("demo_data", exist_ok=True)

header_fill = PatternFill(start_color="7B2CBF", end_color="7B2CBF", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)

def style_sheet(ws):
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = openpyxl.utils.get_column_letter(col)
        ws.column_dimensions[col_letter].width = 28

# =====================================================================
# 1. demo_registry.xlsx — СКВОЗНОЙ ИМПОРТ (Единый реестр взаимодействий)
# Строго 10 колонок по ТЗ (Название учреждения*, Вендор*, ПО* обязательны)
# =====================================================================
wb_registry = openpyxl.Workbook()
ws_reg = wb_registry.active
ws_reg.title = "Реестр"

reg_headers = [
    "Название учреждения",
    "Вендор",
    "ПО",
    "Номер договора",
    "Подписание лицензии",
    "Срок действия лицензии (год)",
    "Статус по передаче",
    "ФИО Менеджера",
    "Ответственные от учреждения",
    "Комментарий"
]
ws_reg.append(reg_headers)

reg_rows = [
    ["МГТУ им. Н. Э. Баумана", "ООО «Ростелеком-Солар»", "Solar Dozor", "РТК-2026/08-Д", "да", 2027, "Передано учреждению", "Иван Иванов", "Сергеева М. В.", "Внедрение ПО завершено"],
    ["Университет ИТМО", "ПАО «Ростелеком»", "РТК-Платформа", "РТК-2026/14-Д", "да", 2028, "Передано учреждению", "Андрей Махт", "Ковалев Д. И.", "Обучение преподавателей"],
    ["НИЯУ МИФИ", "ООО «Ростелеком-Солар»", "Solar Dozor", "РТК-2026/06-Д", "нет", 2026, "В процессе передачи", "Павел Милючихин", "Воронов А. П.", "Согласование договора"],
    ["Университет Иннополис", "ПАО «Ростелеком»", "РТК-Платформа", "", "нет", "", "Не передавалось", "Иван Иванов", "Ситникова Е. Б.", "Первичные переговоры"],
    ["СПбПУ Петра Великого", "ПАО «Ростелеком»", "РТК-Инфраструктура", "РТК-2026/07-Д", "да", 2027, "Передано учреждению", "Андрей Махт", "Смирнов О. А.", "Курс запущен"],
    ["НИУ ВШЭ", "ПАО «Ростелеком»", "РТК-Платформа", "РТК-2026/11-Д", "да", 2026, "Передано учреждению", "Павел Милючихин", "Лебедева А. Ю.", "Практические занятия"]
]
for r in reg_rows:
    ws_reg.append(r)
style_sheet(ws_reg)
wb_registry.save("demo_data/demo_registry.xlsx")

# =====================================================================
# 2. demo_students.xlsx — СПИСКИ СТУДЕНТОВ
# =====================================================================
wb_stud = openpyxl.Workbook()
ws_stud = wb_stud.active
ws_stud.title = "Студенты"
ws_stud.append(["ФИО", "Email"])
stud_rows = [
    ["Алексеев Дмитрий Сергеевич", "d.alekseev@edu.bmstu.ru"],
    ["Борисова Екатерина Андреевна", "e.borisova@edu.bmstu.ru"],
    ["Васильев Максим Игоревич", "m.vasiliev@edu.bmstu.ru"],
    ["Гордеева Полина Викторовна", "p.gordeeva@edu.itmo.ru"],
    ["Данилов Артем Романович", "a.danilov@edu.itmo.ru"],
    ["Ермакова София Михайловна", "s.ermakova@edu.mephi.ru"]
]
for r in stud_rows:
    ws_stud.append(r)
style_sheet(ws_stud)
wb_stud.save("demo_data/demo_students.xlsx")

# =====================================================================
# 3. demo_institutions.xlsx — СПИСОК УЧРЕЖДЕНИЙ
# =====================================================================
wb_inst = openpyxl.Workbook()
ws_inst = wb_inst.active
ws_inst.title = "Учреждения"
ws_inst.append(["Наименование организации", "Регион"])
inst_rows = [
    ["МГТУ им. Н. Э. Баумана", "Москва"],
    ["Университет ИТМО", "Санкт-Петербург"],
    ["НИЯУ МИФИ", "Москва"],
    ["Университет Иннополис", "Татарстан"],
    ["СПбПУ Петра Великого", "Санкт-Петербург"],
    ["НИУ ВШЭ", "Москва"]
]
for r in inst_rows:
    ws_inst.append(r)
style_sheet(ws_inst)
wb_inst.save("demo_data/demo_institutions.xlsx")

# =====================================================================
# 4. demo_products.xlsx — ИТ-ПРОДУКТЫ
# =====================================================================
wb_prod = openpyxl.Workbook()
ws_prod = wb_prod.active
ws_prod.title = "Продукты"
ws_prod.append(["Наименование ПО", "Вендор", "Номер договора", "Подписание лицензии", "Срок действия", "Статус передачи"])
prod_rows = [
    ["Solar Dozor", "ООО «Ростелеком-Солар»", "РТК-2026/01", "да", 2027, "Передано учреждению"],
    ["РТК-Платформа", "ПАО «Ростелеком»", "РТК-2026/02", "да", 2028, "Передано учреждению"],
    ["РТК-Инфраструктура", "ПАО «Ростелеком»", "", "нет", "", "Не передавалось"]
]
for r in prod_rows:
    ws_prod.append(r)
style_sheet(ws_prod)
wb_prod.save("demo_data/demo_products.xlsx")

# =====================================================================
# 5. demo_directions.xlsx — ИТ-НАПРАВЛЕНИЯ
# =====================================================================
wb_dir = openpyxl.Workbook()
ws_dir = wb_dir.active
ws_dir.title = "Направления"
ws_dir.append(["Направление", "Название программы"])
dir_rows = [
    ["Информационная безопасность", "Кибербезопасность предприятия"],
    ["DevOps", "DevOps практики и CI/CD"],
    ["Облачные технологии", "Администрирование облачной инфраструктуры"],
    ["Data Science", "Машинное обучение и искусственный интеллект"],
    ["QA", "Автоматизированное тестирование ПО"]
]
for r in dir_rows:
    ws_dir.append(r)
style_sheet(ws_dir)
wb_dir.save("demo_data/demo_directions.xlsx")

print("Все демо-файлы сформированы в папке demo_data/!")